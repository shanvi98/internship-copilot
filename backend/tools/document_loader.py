"""
Extract plain text from an uploaded resume / JD file (PDF, DOCX, TXT, MD).

pypdf and python-docx are imported lazily so the rest of the app works without
them; a missing parser produces a clear error instead of an import crash.
"""

import io
import re
from pathlib import Path

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r", "").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as e:
        raise DocumentError("PDF support isn't installed on the server (pip install pypdf)") from e
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("This PDF is password-protected — export an unprotected copy and try again")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as e:
        raise DocumentError(f"Couldn't read this PDF: {e}") from e


def _docx_text(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise DocumentError("DOCX support isn't installed on the server (pip install python-docx)") from e
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:  # python-docx raises a variety of zip/xml errors
        raise DocumentError(f"Couldn't read this DOCX file: {e}") from e
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # many resume templates lay out sections in tables
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError(f"File is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError(f"Unsupported file type '{ext or 'unknown'}' — use PDF, DOCX, TXT or MD")

    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    else:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")

    text = _clean(text)
    if not text:
        raise DocumentError("No text found in this file (a scanned/image-only PDF needs OCR first)")
    return text
